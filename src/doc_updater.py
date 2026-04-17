import os
import json
from docx import Document

def asegurar_texto(valor, es_lista_nombres=False):
    """
    Función de seguridad: Convierte lo que mande Gemini (sea lista, nulo o texto)
    en un String válido para que Word no se rompa.
    """
    if isinstance(valor, list):
        if es_lista_nombres:
            return ", ".join(str(x) for x in valor)
        else:
            return "\n".join(f"• {str(x)}" for x in valor)
    return str(valor) if valor else ""

def actualizar_word_requerimiento(ruta_doc, datos_json, ticket_id, app_nombre, fecha_reunion):
    print(f"📝 Procesando actualización de documento para {ticket_id}...")
    
    if not os.path.exists(ruta_doc):
        print(f"   ❌ No se encontró el documento en: {ruta_doc}")
        return False

    try:
        doc = Document(ruta_doc)
        
        # 1. Parsear el JSON
        try:
            texto_limpio = datos_json.strip()
            if texto_limpio.startswith("```json"):
                texto_limpio = texto_limpio[7:-3].strip()
            elif texto_limpio.startswith("```"):
                texto_limpio = texto_limpio[3:-3].strip()
            data = json.loads(texto_limpio)
        except json.JSONDecodeError as e:
            print(f"   ❌ Error JSON: {e}")
            return False

        # --- 2. Reemplazos Simples (Pasados por el filtro de seguridad) ---
        reemplazos_simples = {
            "{{TICKET}}": str(ticket_id),
            "{{APP}}": str(app_nombre),
            "{{OBJETIVO}}": asegurar_texto(data.get("OBJETIVO", "Objetivo no detectado.")),
            "{{FAQ}}": asegurar_texto(data.get("FAQ", ""))
        }
        
        for p in doc.paragraphs:
            for k, v in reemplazos_simples.items():
                if k in p.text:
                    p.text = p.text.replace(k, v)
                    
        # --- 3. Gestión de Tabla de VERSIONES ---
        tabla_v = None
        for t in doc.tables:
            if t.rows and "Versión" in t.rows[0].cells[0].text:
                tabla_v = t
                break
                
        if tabla_v:
            fila_objetivo = None
            for fila in tabla_v.rows:
                if "{{FECHA_V}}" in fila.cells[1].text:
                    fila_objetivo = fila
                    break
            
            if fila_objetivo:
                fila_v_final = fila_objetivo
                fila_v_final.cells[0].text = "v1.0"
            else:
                ultima_fila = tabla_v.rows[-1]
                v_actual_txt = ultima_fila.cells[0].text.strip()
                try:
                    nro_version = int(v_actual_txt.split(".")[-1]) + 1
                    fila_v_final = tabla_v.add_row()
                    fila_v_final.cells[0].text = f"v1.{nro_version}"
                except ValueError:
                    fila_v_final = tabla_v.add_row()
                    fila_v_final.cells[0].text = "v1.x"
            
            fila_v_final.cells[1].text = str(fecha_reunion)
            fila_v_final.cells[2].text = "Juan Manuel Sandoval"
            fila_v_final.cells[3].text = asegurar_texto(data.get("COMENT_V", "Actualización"))

        # --- 4. Gestión de Tabla de MINUTAS ---
        tabla_m = None
        for t in doc.tables:
            if t.rows and "Minuta" in t.rows[0].cells[0].text:
                tabla_m = t
                break
                
        if tabla_m:
            # Blindamos los participantes
            participantes = asegurar_texto(data.get('PARTICIPANTES', 'Usuario'), es_lista_nombres=True)
            if not participantes or participantes.strip() == "" or participantes.lower() == "none":
                participantes = "Usuario"
                
            reunion_info = f"{fecha_reunion} – Reunión con {participantes}"
            
            fila_m_final = None
            for fila in tabla_m.rows:
                if "{{REUNION_INFO}}" in fila.cells[0].text:
                    fila_m_final = fila
                    break
            
            if not fila_m_final:
                fila_m_final = tabla_m.add_row()

            fila_m_final.cells[0].text = reunion_info
            
            # Blindamos el detalle de la minuta
            fila_m_final.cells[1].text = asegurar_texto(data.get("MINUTA_DETALLE", ""))

        # --- 5. Guardar ---
        doc.save(ruta_doc)
        print("   ✅ ¡Documento actualizado exitosamente!")
        return True

    except Exception as e:
        print(f"   ❌ Error fatal al actualizar el Word: {e}")
        return False